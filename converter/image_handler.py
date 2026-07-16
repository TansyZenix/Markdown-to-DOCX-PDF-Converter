from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt


class ImageHandler:
    def __init__(self, document, styles, source_directory: Path) -> None:
        self.document = document
        self.styles = styles
        self.source_directory = source_directory

    def add_image(self, source: str, alt_text: str, number: int) -> None:
        image_path = (self.source_directory / source).resolve()
        if not image_path.is_file():
            raise FileNotFoundError(f"Image not found: {image_path}")
        section = self.document.sections[0]
        max_width = section.page_width - section.left_margin - section.right_margin
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.add_run().add_picture(str(image_path), width=max_width * self.styles.config["images"]["max_width_ratio"])
        caption = self.document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        text = f"图 {number}" + (f" {alt_text}" if alt_text else "")
        run = caption.add_run(text)
        self.styles.apply_run_font(run, self.styles.config["captions"]["font_size_pt"], True)
        figure_caption = self.styles.config["captions"]["figure"]
        self.styles.format_paragraph(caption, before=figure_caption["space_before_pt"], after=figure_caption["space_after_pt"])
