"""Markdown to DOCX/PDF conversion package."""
import re
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

_original_add_run = Paragraph.add_run

def _font(run, name):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia"): fonts.set(qn(f"w:{slot}"), name)

def _add_run(self, text=None, style=None):
    if text and self.style.name in ("Heading 1", "Heading 2", "Heading 3") and not self.runs:
        state = getattr(self.part, "_md_heading_state", {"second": None, "third": 0}); prefix = re.match(r"^(\d+(?:\.\d+)*)\.?\s*", text)
        if self.style.name == "Heading 2" and prefix: state = {"second": prefix.group(1), "third": 0}
        elif self.style.name == "Heading 3" and not prefix and state["second"]:
            state["third"] += 1; text = f"{state['second']}.{state['third']} {text}"
        self.part._md_heading_state = state
    if not text: return _original_add_run(self, text, style)
    pieces = re.findall(r"[\u4e00-\u9fff]+|[^\u4e00-\u9fff]+", text)
    first = None
    for piece in pieces:
        run = _original_add_run(self, piece, style)
        _font(run, "宋体" if re.fullmatch(r"[\u4e00-\u9fff]+", piece) else "Times New Roman")
        first = first or run
    return first

Paragraph.add_run = _add_run
