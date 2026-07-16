"""DOCX generator with academic heading level mapping."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Cm

from .formula_handler import FormulaHandler
from .image_handler import ImageHandler
from .markdown_parser import MarkdownParser
from .table_handler import TableHandler


class DocxGenerator:
    """Converts Markdown tokens into a python-docx Document."""

    def __init__(self, styles, source_directory: Path) -> None:
        self.styles = styles
        self.source_directory = Path(source_directory)
        self.figure_number = 0
        self.table_number = 0
        self.formulas: list[str] = []

    def convert(self, source: str) -> Document:
        source = self._extract_formulas(source)
        document = Document()
        self.styles.apply_document_defaults(document)
        tokens = self._parse(source)
        image_handler = ImageHandler(document, self.styles, self.source_directory)
        table_handler = TableHandler(document, self.styles)
        list_stack: list[str] = []
        quote_depth = 0
        i = 0
        while i < len(tokens):
            token = tokens[i]
            if token.type in ("heading_open", "paragraph_open"):
                heading_level = min(int(token.tag[1]), 3)
                # Shift heading levels for academic hierarchy: markdown ## → Word Heading 1
                style_name = (
                    f"Heading {heading_level}"
                    if token.type == "heading_open"
                    else "Normal"
                )
                paragraph = document.add_paragraph(style=style_name)
                if token.type == "paragraph_open" and list_stack:
                    paragraph.style = (
                        "List Number" if list_stack[-1] == "o" else "List Bullet"
                    )
                if token.type == "paragraph_open" and quote_depth:
                    paragraph.paragraph_format.left_indent = Cm(
                        self.styles.config["blockquote"]["left_indent_cm"] * quote_depth
                    )
                self._inline(paragraph, tokens[i + 1].children or [], image_handler)
                i += 3
                continue
            if token.type == "bullet_list_open":
                list_stack.append("b")
            elif token.type == "ordered_list_open":
                list_stack.append("o")
            elif token.type in ("bullet_list_close", "ordered_list_close"):
                list_stack.pop()
            elif token.type == "blockquote_open":
                quote_depth += 1
            elif token.type == "blockquote_close":
                quote_depth -= 1
            elif token.type == "fence":
                paragraph = document.add_paragraph(style="Code Block")
                run = paragraph.add_run(token.content.rstrip("\n"))
                self._code(run)
            elif token.type == "table_open":
                i = self._table(tokens, i, table_handler)
                continue
            i += 1
        return document

    def _parse(self, source: str):
        """Parse markdown and shift heading levels for academic hierarchy."""
        tokens = MarkdownParser().parse(source)
        for token in tokens:
            if token.type == "heading_open" and token.tag in ("h2", "h3", "h4"):
                token.tag = f"h{int(token.tag[1]) - 1}"
        return tokens

    def _extract_formulas(self, source: str) -> str:
        """Extract LaTeX formulas and replace with placeholders."""
        def put(match):
            self.formulas.append(match.group(1))
            return f"\n[[MATH:{len(self.formulas) - 1}]]\n"

        source = re.sub(r"\$\$([\s\S]+?)\$\$", put, source)
        return re.sub(r"(?<!\$)\$([^$\n]+)\$(?!\$)", put, source)

    def _inline(self, paragraph, children, image_handler) -> None:
        bold = italic = False
        for child in children:
            if child.type == "text":
                self._text(paragraph, child.content, bold, italic)
            elif child.type == "code_inline":
                run = paragraph.add_run(child.content)
                self._code(run)
                run.bold = bold
                run.italic = italic
            elif child.type in ("softbreak", "hardbreak"):
                paragraph.add_run().add_break()
            elif child.type == "strong_open":
                bold = True
            elif child.type == "strong_close":
                bold = False
            elif child.type == "em_open":
                italic = True
            elif child.type == "em_close":
                italic = False
            elif child.type == "image":
                self.figure_number += 1
                image_handler.add_image(
                    child.attrGet("src"), child.content, self.figure_number
                )

    def _text(self, paragraph, text: str, bold: bool, italic: bool) -> None:
        for segment in re.split(
            r"(\[\[MATH:\d+\]\]|\*\*.+?\*\*|`[^`]+`)", text
        ):
            if not segment:
                continue
            math_match = re.fullmatch(r"\[\[MATH:(\d+)\]\]", segment)
            if math_match:
                FormulaHandler().add(paragraph, self.formulas[int(math_match.group(1))])
                continue
            if segment.startswith("**"):
                run = paragraph.add_run(segment[2:-2])
                self.styles.apply_run_font(run)
                run.bold = True
                run.italic = italic
            elif segment.startswith("`"):
                run = paragraph.add_run(segment[1:-1])
                self._code(run)
                run.bold = bold
                run.italic = italic
            else:
                # Remove extra space between CJK and Latin characters
                cleaned = re.sub(
                    r"(?<=[一-鿿]) +(?=[A-Za-z0-9])|(?<=[A-Za-z0-9]) +(?=[一-鿿])",
                    "",
                    segment,
                )
                run = paragraph.add_run(cleaned)
                self.styles.apply_run_font(run)
                run.bold = bold
                run.italic = italic

    def _code(self, run) -> None:
        self.styles.apply_run_font(run, self.styles.config["code"]["font_size_pt"])
        code_font = self.styles.config["fonts"]["code"]
        run.font.name = code_font
        from docx.oxml.ns import qn

        run_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for slot in ("ascii", "hAnsi", "eastAsia"):
            run_fonts.set(qn(f"w:{slot}"), code_font)

    def _table(self, tokens, start, table_handler) -> int:
        header = []
        rows = []
        row = []
        in_header = False
        i = start + 1
        while i < len(tokens) and tokens[i].type != "table_close":
            token = tokens[i]
            if token.type == "thead_open":
                in_header = True
            elif token.type == "thead_close":
                in_header = False
            elif token.type == "tr_open":
                row = []
            elif token.type in ("th_open", "td_open"):
                row.append(tokens[i + 1].content)
            elif token.type == "tr_close":
                if in_header:
                    header = row
                else:
                    rows.append(row)
            i += 1
        self.table_number += 1
        table_handler.add_table(header, rows, self.table_number)
        return i + 1
