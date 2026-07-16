from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .markdown_parser import MarkdownParser


def _set_border(table, edge: str, size_eighths: int | None) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    element = borders.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        borders.append(element)
    element.set(qn("w:val"), "single" if size_eighths else "nil")
    if size_eighths:
        element.set(qn("w:sz"), str(size_eighths))
        element.set(qn("w:color"), "000000")


class TableHandler:
    def __init__(self, document, styles) -> None:
        self.document, self.styles, self.parser = document, styles, MarkdownParser()

    def add_table(self, headers: list[str], rows: list[list[str]], number: int) -> None:
        caption = self.document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(f"表 {number} " + " ".join(headers[:1]))
        self.styles.apply_run_font(run, self.styles.config["captions"]["font_size_pt"], True)
        settings = self.styles.config["captions"]["table"]
        self.styles.format_paragraph(caption, before=settings["space_before_pt"], after=settings["space_after_pt"])
        table = self.document.add_table(rows=1, cols=len(headers))
        for index, value in enumerate(headers):
            self._write_cell(table.rows[0].cells[index], value, bold=True)
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                self._write_cell(cells[index], value)
        for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
            _set_border(table, edge, None)
        border = self.styles.config["tables"]
        _set_border(table, "top", round(border["top_bottom_border_pt"] * 8))
        _set_border(table, "bottom", round(border["top_bottom_border_pt"] * 8))
        for cell in table.rows[0].cells:
            borders = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcBorders")
            if borders is None:
                borders = OxmlElement("w:tcBorders")
                cell._tc.get_or_add_tcPr().append(borders)
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), str(round(border["header_border_pt"] * 8)))
            bottom.set(qn("w:color"), "000000")
            borders.append(bottom)

    def _write_cell(self, cell, value: str, bold: bool = False) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.styles.format_paragraph(paragraph)
        active_bold, active_italic = bold, False
        for token in self.parser.parse_inline(value):
            if token.type == "text":
                run = paragraph.add_run(token.content)
                self.styles.apply_run_font(run, self.styles.config["tables"]["font_size_pt"])
                run.bold, run.italic = active_bold, active_italic
            elif token.type == "code_inline":
                run = paragraph.add_run(token.content)
                self.styles.apply_run_font(run, self.styles.config["tables"]["font_size_pt"])
                code_font = self.styles.config["fonts"]["code"]
                run.font.name = code_font
                rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                for slot in ("ascii", "hAnsi", "eastAsia"):
                    rfonts.set(qn(f"w:{slot}"), code_font)
                run.bold, run.italic = active_bold, active_italic
            elif token.type == "strong_open":
                active_bold = True
            elif token.type == "strong_close":
                active_bold = bold
            elif token.type == "em_open":
                active_italic = True
            elif token.type == "em_close":
                active_italic = False
