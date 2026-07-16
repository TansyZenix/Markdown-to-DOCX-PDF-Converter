from pathlib import Path

from docx import Document
import pytest

from converter.style_manager import StyleManager
from converter.environment import check_environment
from converter.markdown_parser import MarkdownParser
from converter.docx_generator import DocxGenerator
from converter.pdf_exporter import LibreOfficeNotFoundError, build_pdf_command, export_pdf
from main import resolve_output_path


ROOT = Path(__file__).resolve().parents[1]


def test_style_manager_applies_a4_fonts_and_heading_styles():
    document = Document()
    manager = StyleManager.load(ROOT / "config" / "style.yaml")
    manager.apply_document_defaults(document)
    assert document.sections[0].page_width.cm == pytest.approx(21.0, abs=0.01)
    assert document.styles["Normal"].font.size.pt == 10.5
    assert document.styles["Heading 1"].font.size.pt > document.styles["Heading 2"].font.size.pt
    run = document.add_paragraph().add_run("中文 English")
    manager.apply_run_font(run)
    assert run.font.name == "Times New Roman"


def test_environment_check_warns_for_unexpected_conda_name(monkeypatch):
    monkeypatch.setenv("CONDA_DEFAULT_ENV", "otherEnv")
    report = check_environment(required_environment="md2docx-env", required_modules=())
    assert report.warnings and "md2docx-env" in report.warnings[0]


def test_parser_emits_tokens_for_required_markdown_blocks():
    source = "# Title\n\n**bold** and *italic*\n\n- item\n\n1. first\n\n> quote\n\n```python\nprint('ok')\n```\n\n| A | B |\n| - | - |\n| 1 | 2 |\n"
    token_types = [token.type for token in MarkdownParser().parse(source)]
    for expected in ("heading_open", "bullet_list_open", "ordered_list_open", "blockquote_open", "fence", "table_open"):
        assert expected in token_types


def test_generator_maps_headings_and_inline_emphasis():
    generator = DocxGenerator(StyleManager.load(ROOT / "config" / "style.yaml"), ROOT)
    document = generator.convert("# 一级标题\n\n正文含 **加粗** 与 *斜体*。")
    assert document.paragraphs[0].style.name == "Heading 1"
    assert any(run.bold for run in document.paragraphs[1].runs)
    assert any(run.italic for run in document.paragraphs[1].runs)


def test_generator_creates_native_three_line_table_and_figure_caption():
    generator = DocxGenerator(StyleManager.load(ROOT / "config" / "style.yaml"), ROOT / "examples")
    document = generator.convert("| A | B |\n| - | - |\n| 1 | 2 |\n\n![示例](../image.png)")
    assert len(document.tables) == 1
    assert document.paragraphs[0].text.startswith("表 1")
    assert any(paragraph.text.startswith("图 1") for paragraph in document.paragraphs)
    assert document.inline_shapes


def test_pdf_command_targets_docx_and_output_directory(tmp_path):
    command = build_pdf_command(Path("C:/LibreOffice/soffice.exe"), tmp_path / "sample.docx", tmp_path)
    assert "--headless" in command and str(tmp_path) in command and command[-1].endswith("sample.docx")


def test_export_pdf_reports_missing_executable(tmp_path):
    with pytest.raises(LibreOfficeNotFoundError):
        export_pdf(tmp_path / "input.docx", tmp_path, executable="missing-soffice.exe")


def test_cli_defaults_output_to_input_stem(tmp_path):
    assert resolve_output_path(tmp_path / "input.md", None) == tmp_path / "input.docx"
