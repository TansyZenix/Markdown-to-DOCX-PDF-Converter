from pathlib import Path

from docx import Document

from converter.docx_generator import DocxGenerator
from converter.style_manager import StyleManager


ROOT = Path(__file__).resolve().parents[1]


def test_demo_markdown_converts_to_editable_docx_with_native_objects(tmp_path):
    source = ROOT / "examples" / "demo.md"
    document = DocxGenerator(StyleManager.load(ROOT / "config" / "style.yaml"), source.parent).convert(
        source.read_text(encoding="utf-8")
    )
    output = tmp_path / "demo.docx"
    document.save(output)
    rendered = Document(output)

    assert rendered.paragraphs[0].style.name == "Heading 1"
    assert any(paragraph.style.name == "Heading 2" for paragraph in rendered.paragraphs)
    assert len(rendered.tables) == 1
    assert len(rendered.inline_shapes) == 1
    assert rendered.styles["Normal"].font.size.pt == 10.5
